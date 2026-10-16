"""Convert docs/phase1_report.md into a formatted Word document.

    python docs/build_report.py [--results results/phase1-baseline-local]

Supports the Markdown subset used in the report: headings, paragraphs, **bold**,
*italic*, bullet and numbered lists, pipe tables, images and <!-- pagebreak -->.
``{{INITIAL_RESULTS}}`` is replaced with tables generated from a results directory.
"""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT  # noqa: F401  (kept for manual landscape tweaks)
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

DOCS = Path(__file__).resolve().parent
ROOT = DOCS.parent
INLINE = re.compile(r"(\*\*.+?\*\*|\*.+?\*|`.+?`)")


def add_inline(par, text: str, size: float | None = None, bold: bool = False) -> None:
    text = text.replace("\\*", "*")
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            add_inline(par, part[2:-2], size=size, bold=True)
            continue
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = par.add_run(part[1:-1])
            run.italic = True
            run.bold = bold
        elif part.startswith("`") and part.endswith("`"):
            run = par.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.bold = bold
        else:
            run = par.add_run(part)
            run.bold = bold
        if size:
            run.font.size = Pt(size)


def shade(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def add_table(doc, rows: list[list[str]]) -> None:
    ncols = len(rows[0])
    size = 11 if ncols <= 3 else (9 if ncols <= 5 else 7.5)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    # column widths proportional to (capped) average text length
    weights = []
    for j in range(ncols):
        cells = [r[j] for r in rows if j < len(r)]
        avg = sum(len(c) for c in cells) / len(cells)
        longest_word = max((len(w) for c in cells for w in c.replace("*", "").split()), default=1)
        weights.append(min(max(avg, longest_word * 1.3, 3), 60))
    total_w = 17.0
    widths = [Cm(total_w * w / sum(weights)) for w in weights]
    for i, row in enumerate(rows):
        for j, cell_text in enumerate(row):
            cell = table.cell(i, j)
            cell.width = widths[j]
            cell.text = ""
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(0)
            add_inline(par, cell_text.strip(), size=size, bold=(i == 0))
            if i == 0:
                shade(cell, "D9E2F3")
            elif "BASE" in row[1] if len(row) > 1 else False:
                shade(cell, "FFF2CC")
    doc.add_paragraph()


def results_markdown(results_dir: Path) -> str:
    path = results_dir / "episodes.jsonl"
    if not path.exists():
        return "*Initial results: run `python -m mcpguard.eval.run --config configs/baseline_local.yaml`.*"
    df = pd.DataFrame([json.loads(l) for l in path.read_text(encoding="utf-8").splitlines()])
    lines = ["**Initial results (Phase 1 baseline).** Every attacked and benign scenario was run "
             "end to end through real MCP servers; numbers are percentages.", ""]
    lines.append("| Model | Defense | Episodes | BU | ASR | PUA | NRP | PVR | Tool calls/ep | Sec/ep |")
    lines.append("|---|---|---|---|---|---|---|---|---|---|")
    for (model, defense), g in df.groupby(["model", "defense"], sort=False):
        att, ben = g[g.attack_id.notna()], g[g.attack_id.isna()]
        asr, pua = att.attack_success.mean(), att.utility.mean()
        lines.append(
            f"| {model} | {defense} | {len(g)} | {100*ben.utility.mean():.1f} | {100*asr:.1f} | "
            f"{100*pua:.1f} | {100*pua*(1-asr):.1f} | {100*(att.privilege_violations > 0).mean():.1f} | "
            f"{att.tool_calls.mean():.2f} | {g.wall_time_s.mean():.1f} |")
    lines += ["", "**ASR (%) by attack class.**", ""]
    att = df[df.attack_id.notna()]
    piv = att.pivot_table(index="attack_class", columns=["model", "defense"], values="attack_success",
                          aggfunc="mean") * 100
    order = [c for c in ["PI", "NC", "PM", "OP", "UI", "FE", "TT", "RI", "PE", "CD"] if c in piv.index]
    piv = piv.loc[order]
    header = ["Class"] + [f"{m.split(':', 1)[1]} / {d}" for m, d in piv.columns]
    lines.append("| " + " | ".join(header) + " |")
    lines.append("|" + "---|" * len(header))
    for cls, row in piv.iterrows():
        lines.append("| " + " | ".join([cls] + [f"{v:.0f}" for v in row.values]) + " |")
    notes = DOCS / "phase1_results_notes.md"
    if notes.exists():
        lines += ["", notes.read_text(encoding="utf-8").strip()]
    return "\n".join(lines)


def build(md_path: Path, out_path: Path, results_dirs: list[Path]) -> None:
    text = md_path.read_text(encoding="utf-8")
    text = text.replace("{{INITIAL_RESULTS}}", "\n\n".join(results_markdown(d) for d in results_dirs))
    doc = Document()
    for section in doc.sections:
        section.left_margin = section.right_margin = Cm(2.0)
        section.top_margin = section.bottom_margin = Cm(2.0)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    for lvl, size in ((1, 15), (2, 13), (3, 12)):
        st = doc.styles[f"Heading {lvl}"]
        st.font.name = "Times New Roman"
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
        rfonts = st.element.get_or_add_rPr().get_or_add_rFonts()
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rfonts.set(qn(attr), "Times New Roman")
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            rfonts.attrib.pop(qn(attr), None)

    lines = text.splitlines()
    i = 0
    first_heading = True
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        if line.strip() == "<!-- pagebreak -->":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            i += 1
            continue
        m = re.match(r"^(#{1,3})\s+(.*)", line)
        if m:
            level = len(m.group(1))
            if first_heading:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(m.group(2))
                r.bold = True
                r.font.size = Pt(18)
                first_heading = False
            else:
                doc.add_heading(m.group(2), level=level)
            i += 1
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r"\s*:?-{2,}:?\s*", c) for c in cells):
                    rows.append(cells)
                i += 1
            add_table(doc, rows)
            continue
        m = re.match(r"^!\[(.*?)\]\((.*?)\)", line)
        if m:
            doc.add_picture(str((DOCS / m.group(2)).resolve()), width=Cm(17))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue
        m = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)", line)
        if m:
            indent = len(m.group(1)) >= 2
            if m.group(2) in "-*":
                p = doc.add_paragraph(style="List Bullet 2" if indent else "List Bullet")
            else:  # explicit numbers so each list restarts at its own value
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(1.6 if indent else 0.9)
                p.paragraph_format.first_line_indent = Cm(-0.6)
                p.paragraph_format.space_after = Pt(2)
                p.add_run(m.group(2) + " ")
            add_inline(p, m.group(3))
            i += 1
            continue
        # paragraph: merge consecutive non-empty plain lines
        buf = [line.strip()]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"^(#|\||!\[|\s*[-*]\s|\s*\d+\.\s|<!--)", lines[i]):
            buf.append(lines[i].strip())
            i += 1
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        joined = "\n".join(buf)
        for k, chunk in enumerate(joined.split("\n")):
            if k:
                p.add_run().add_break()
            add_inline(p, chunk)
    doc.save(out_path)
    print(f"saved {out_path}")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--results", nargs="*", default=[str(ROOT / "results" / "phase1-baseline-local")])
    ap.add_argument("--out", default=str(DOCS / "Phase1_G1-XX_Report.docx"))
    a = ap.parse_args()
    build(DOCS / "phase1_report.md", Path(a.out), [Path(r) for r in a.results])
